"""The roster spreadsheet: the template we hand out, and the one we get back.

The exchange rests on the player id, so these check what each kind of row means:
an id with a different name is a rename, no id is somebody new, and a missing row
is someone who left.
"""
from __future__ import annotations

import io
from datetime import date

import pytest
import pytest_asyncio
from fastapi import FastAPI, HTTPException, status
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.ext.compiler import compiles

from dwsbot import roster_sheet
from dwsbot.api.deps import current_user, get_session, require_admin
from dwsbot.api.routers import players as players_router
from dwsbot.db import Base
from dwsbot.models import AuditLog, Player, PlayerName
from dwsbot.schemas import MeOut


@compiles(JSONB, "sqlite")
def _jsonb_on_sqlite(type_, compiler, **kw):  # noqa: ARG001
    return "JSON"


ADMIN = MeOut(discord_id=1, username="Admin", is_admin=True)
MEMBER = MeOut(discord_id=2, username="Member", is_admin=False)
DAY = date(2026, 9, 25)


@pytest_asyncio.fixture
async def client_factory():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    maker = async_sessionmaker(engine, expire_on_commit=False)

    def build(user: MeOut = ADMIN):
        app = FastAPI()
        app.include_router(players_router.router)

        async def _session():
            async with maker() as s:
                yield s

        async def _admin():
            if not user.is_admin:
                raise HTTPException(status.HTTP_403_FORBIDDEN, "admins only")
            return user

        app.dependency_overrides[get_session] = _session
        app.dependency_overrides[current_user] = lambda: user
        app.dependency_overrides[require_admin] = _admin
        return AsyncClient(transport=ASGITransport(app=app), base_url="http://t")

    build.maker = maker
    yield build
    await engine.dispose()


async def seed(maker, *people) -> dict[str, str]:
    """Players in the database, returning name -> id."""
    ids = {}
    async with maker() as session:
        for name, fields in people:
            player = Player(name=name, **fields)
            player.names.append(PlayerName(name=name, first_seen=date(2026, 9, 9),
                                           last_seen=date(2026, 9, 9)))
            session.add(player)
            await session.flush()
            ids[name] = str(player.id)
        await session.commit()
    return ids


def sheet_of(rows: list[dict]) -> bytes:
    """A workbook in the template's shape, as someone would upload it."""
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = roster_sheet.SHEET
    sheet.append([heading for heading, _ in roster_sheet.COLUMNS])
    for row in rows:
        sheet.append([row.get(attribute) for _, attribute in roster_sheet.COLUMNS])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


async def upload(client, data: bytes, as_of: date = DAY):
    return await client.post("/players/import/preview",
                             files={"file": ("roster.xlsx", data, "application/vnd.ms-excel")},
                             data={"as_of": as_of.isoformat()})


@pytest.mark.asyncio
async def test_only_admins_may_take_the_template(client_factory):
    async with client_factory(MEMBER) as c:
        assert (await c.get("/players/template.xlsx")).status_code == 403


@pytest.mark.asyncio
async def test_the_template_carries_every_current_member_and_their_id(client_factory):
    ids = await seed(client_factory.maker,
                     ("Dubai88", {"rank": 4, "industry_level": 8, "bgb_cp": 292_185_930,
                                  "total_cp": 1_694_095_810}),
                     ("Ck360", {"rank": 3, "bgb_cp": 192_063_851}),
                     ("Gone", {"active": False}))
    async with client_factory() as c:
        r = await c.get("/players/template.xlsx")
    assert r.status_code == 200
    assert "pou-roster-" in r.headers["content-disposition"]
    book = load_workbook(io.BytesIO(r.content))
    sheet = book[roster_sheet.SHEET]
    assert [c.value for c in sheet[1]] == [h for h, _ in roster_sheet.COLUMNS]
    body = {row[1]: row for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert set(body) == {"Dubai88", "Ck360"}          # someone who left is not offered
    assert body["Dubai88"][0] == ids["Dubai88"]       # the id comes along
    assert body["Dubai88"][5] == 1_694_095_810        # numbers, not text
    assert roster_sheet.HELP_SHEET in book.sheetnames


@pytest.mark.asyncio
async def test_preview_reads_each_kind_of_row_and_writes_nothing(client_factory):
    ids = await seed(client_factory.maker,
                     ("Dubai88", {"bgb_cp": 292_185_930, "rank": 4}),
                     ("SandySQ", {"bgb_cp": 35_486_818}),
                     ("Leaver", {"bgb_cp": 1_000_000}))
    data = sheet_of([
        {"id": ids["Dubai88"], "name": "Dubai88", "bgb_cp": "300,000,000", "rank": 4},
        {"id": ids["SandySQ"], "name": "Sandy", "bgb_cp": 35_486_818},
        {"name": "Hoshiko", "bgb_cp": 18_010_756, "industry_level": 4},
    ])
    async with client_factory() as c:
        r = await upload(c, data)
        assert r.status_code == 200, r.text
        body = r.json()
        listed = (await c.get("/players")).json()
    assert body["problems"] == []
    assert [c["name"] for c in body["updated"]] == ["Dubai88"]
    assert body["updated"][0]["fields"]["bgb_cp"] == [292_185_930, 300_000_000]  # commas read
    assert [(c["was"], c["name"]) for c in body["renamed"]] == [("SandySQ", "Sandy")]
    assert [c["name"] for c in body["added"]] == ["Hoshiko"]
    assert [c["name"] for c in body["left"]] == ["Leaver"]
    assert {p["name"] for p in listed} == {"Dubai88", "SandySQ", "Leaver"}   # nothing written yet


@pytest.mark.asyncio
async def test_applying_writes_exactly_what_was_shown(client_factory):
    ids = await seed(client_factory.maker,
                     ("Dubai88", {"bgb_cp": 292_185_930}),
                     ("SandySQ", {"bgb_cp": 35_486_818}),
                     ("Leaver", {"bgb_cp": 1_000_000}))
    data = sheet_of([
        {"id": ids["Dubai88"], "name": "Dubai88", "bgb_cp": 300_000_000},
        {"id": ids["SandySQ"], "name": "Sandy", "bgb_cp": 36_000_000},
        {"name": "Hoshiko", "bgb_cp": 18_010_756},
    ])
    async with client_factory() as c:
        preview = (await upload(c, data)).json()
        r = await c.post("/players/import/apply", json={
            "as_of": DAY.isoformat(), "fingerprint": preview["fingerprint"],
            "file_name": "roster.xlsx", "rows": preview["rows"], "keep": []})
        assert r.status_code == 200, r.text
        listed = {p["name"]: p for p in (await c.get("/players")).json()}
    assert listed["Dubai88"]["bgb_cp"] == 300_000_000
    assert listed["Sandy"]["active"] is True
    assert [(n["name"], n["last_seen"]) for n in listed["Sandy"]["names"]] == [
        ("SandySQ", "2026-09-09"), ("Sandy", "2026-09-25")]     # the old name is kept
    assert listed["Hoshiko"]["names"][0]["first_seen"] == "2026-09-25"
    assert listed["Leaver"]["active"] is False                   # missing row means gone
    assert [n["last_seen"] for n in listed["Dubai88"]["names"]] == ["2026-09-25"]
    async with client_factory.maker() as s:
        entry = (await s.scalars(select(AuditLog).where(AuditLog.action == "player.import"))).one()
    assert entry.detail["as_of"] == "2026-09-25"
    assert entry.detail["file"] == "roster.xlsx"
    assert entry.detail["uploaded_at"].startswith("20")          # when it was uploaded
    assert entry.detail["counts"] == {"updated": 1, "renamed": 1, "added": 1, "left": 1,
                                      "returning": 0, "unchanged": 0}


@pytest.mark.asyncio
async def test_a_member_can_be_kept_rather_than_marked_as_having_left(client_factory):
    ids = await seed(client_factory.maker, ("Dubai88", {}), ("Missing", {}))
    data = sheet_of([{"id": ids["Dubai88"], "name": "Dubai88"}])
    async with client_factory() as c:
        preview = (await upload(c, data)).json()
        assert [c["name"] for c in preview["left"]] == ["Missing"]
        await c.post("/players/import/apply", json={
            "as_of": DAY.isoformat(), "fingerprint": preview["fingerprint"],
            "rows": preview["rows"], "keep": [ids["Missing"]]})
        listed = {p["name"]: p for p in (await c.get("/players")).json()}
    assert listed["Missing"]["active"] is True


@pytest.mark.asyncio
async def test_someone_who_left_and_came_back(client_factory):
    ids = await seed(client_factory.maker, ("Stinkycät", {"active": False, "bgb_cp": 18_948_867}))
    data = sheet_of([{"id": ids["Stinkycät"], "name": "Stinkycät", "bgb_cp": 20_000_000}])
    async with client_factory() as c:
        preview = (await upload(c, data)).json()
        assert [c["name"] for c in preview["returning"]] == ["Stinkycät"]
        await c.post("/players/import/apply", json={
            "as_of": DAY.isoformat(), "fingerprint": preview["fingerprint"],
            "rows": preview["rows"], "keep": []})
        listed = {p["name"]: p for p in (await c.get("/players")).json()}
    assert listed["Stinkycät"]["active"] is True
    assert listed["Stinkycät"]["bgb_cp"] == 20_000_000


@pytest.mark.asyncio
async def test_a_roster_edited_meanwhile_is_refused(client_factory):
    ids = await seed(client_factory.maker, ("Dubai88", {"bgb_cp": 1}))
    data = sheet_of([{"id": ids["Dubai88"], "name": "Dubai88", "bgb_cp": 2}])
    async with client_factory() as c:
        preview = (await upload(c, data)).json()
        await c.patch(f"/players/{ids['Dubai88']}", json={"rank": 5})   # someone edits meanwhile
        r = await c.post("/players/import/apply", json={
            "as_of": DAY.isoformat(), "fingerprint": preview["fingerprint"],
            "rows": preview["rows"], "keep": []})
    assert r.status_code == 409
    assert "changed" in r.json()["detail"]


@pytest.mark.asyncio
async def test_a_sheet_that_cannot_be_trusted_is_reported_not_guessed(client_factory):
    ids = await seed(client_factory.maker, ("Dubai88", {}))
    data = sheet_of([
        {"id": ids["Dubai88"], "name": "Dubai88"},
        {"id": "not-an-id", "name": "Broken"},
        {"name": "", "bgb_cp": 5},
        {"name": "Twin"}, {"name": "Twin"},
        {"id": "3f7c4b66-0000-4000-8000-000000000000", "name": "Ghost"},
        {"name": "Odd", "rank": 9},
    ])
    async with client_factory() as c:
        body = (await upload(c, data)).json()
    joined = " ".join(body["problems"])
    assert "not a player id" in joined          # a mangled id
    assert "no name" in joined                  # a row with nothing in the name
    assert "cannot share a name" in joined      # two rows called the same
    assert "no player has the id" in joined     # an id that is not ours
    assert "out of range" in joined             # rank 9
    assert body["added"] == [] or all(c["name"] != "Broken" for c in body["added"])


@pytest.mark.asyncio
async def test_an_empty_sheet_does_not_empty_the_roster(client_factory):
    await seed(client_factory.maker, ("Dubai88", {}), ("Ck360", {}))
    async with client_factory() as c:
        preview = (await upload(c, sheet_of([]))).json()
        assert preview["left"] == []
        assert "No rows could be read" in " ".join(preview["problems"])
        r = await c.post("/players/import/apply", json={
            "as_of": DAY.isoformat(), "fingerprint": preview["fingerprint"],
            "rows": [], "keep": []})
        assert r.status_code == 422
        still = {p["name"] for p in (await c.get("/players")).json() if p["active"]}
        assert still == {"Dubai88", "Ck360"}


@pytest.mark.asyncio
async def test_a_file_that_is_not_a_spreadsheet(client_factory):
    await seed(client_factory.maker, ("Dubai88", {}))
    async with client_factory() as c:
        body = (await upload(c, b"this is not a workbook")).json()
    assert "could not be opened as a spreadsheet" in body["problems"][0]
    assert body["left"] == []                   # and it does not read that as everyone leaving
