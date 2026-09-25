"""The roster as a spreadsheet: the template we hand out, and the one we get back.

The exchange is built on the player id. The template carries every current
member's id, so a row that comes back with the same id is that player however
much the name has changed, a row with no id is somebody new, and a member whose
row is gone has left. That makes a rename unmistakable, which is the one thing
reading names off screenshots can never be sure of.

Nothing here touches the database: it produces a plan, and the endpoint applies
it once the person uploading has agreed to it.
"""
from __future__ import annotations

import hashlib
import io
import uuid
from dataclasses import dataclass, field
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SHEET = "Members"
HELP_SHEET = "How to use"

# The column headings, and the field each one fills. Order is what the template
# is written in; an uploaded file is read by heading, so columns may be moved.
COLUMNS: list[tuple[str, str]] = [
    ("Player id", "id"),
    ("Name", "name"),
    ("Rank", "rank"),
    ("Industry level", "industry_level"),
    ("BGB CP", "bgb_cp"),
    ("Total CP", "total_cp"),
]
NUMERIC = {"rank": (1, 5), "industry_level": (1, 99), "bgb_cp": (0, 10**12),
           "total_cp": (0, 10**12)}
WIDTHS = {"Player id": 38, "Name": 24, "Rank": 8, "Industry level": 15,
          "BGB CP": 16, "Total CP": 16}

HELP = [
    "How to update the roster",
    "",
    "1. This file lists every current member, with the id the site knows them by.",
    "2. Edit the numbers and names on the Members sheet, then upload it on the",
    "   Members tab. The site shows what would change and asks you to confirm.",
    "",
    "The Player id column is what makes this reliable:",
    "",
    "  * Leave a row's id alone. It says which player the row is about, so you can",
    "    rename anyone simply by typing the new name.",
    "  * A new member: add a row and leave Player id empty.",
    "  * Someone who left: delete their row.",
    "",
    "Rank is 1 to 5 (R1-R5). Leave a cell empty to leave that value as it is.",
    "Numbers may be typed with or without commas.",
]


@dataclass
class SheetRow:
    """One row as it came back, already checked."""

    line: int
    id: uuid.UUID | None
    name: str
    rank: int | None = None
    industry_level: int | None = None
    bgb_cp: int | None = None
    total_cp: int | None = None


@dataclass
class Change:
    player_id: str | None
    name: str
    was: str | None = None                 # the old name, when it changed
    fields: dict[str, tuple] = field(default_factory=dict)  # field -> (before, after)
    line: int = 0


@dataclass
class Plan:
    as_of: date
    updated: list[Change] = field(default_factory=list)
    renamed: list[Change] = field(default_factory=list)
    added: list[Change] = field(default_factory=list)
    left: list[Change] = field(default_factory=list)
    returning: list[Change] = field(default_factory=list)
    unchanged: int = 0
    problems: list[str] = field(default_factory=list)

    @property
    def writes(self) -> int:
        return len(self.updated) + len(self.renamed) + len(self.added) + \
            len(self.left) + len(self.returning)


def fingerprint(players: list) -> str:
    """A short digest of the roster, to notice it changing under an upload."""
    parts = sorted(
        f"{p.id}|{p.name}|{p.active}|{p.rank}|{p.industry_level}|{p.bgb_cp}|{p.total_cp}"
        for p in players
    )
    return hashlib.sha256("\n".join(parts).encode()).hexdigest()[:32]


def build_template(players: list, as_of: date | None = None) -> bytes:
    """The workbook to hand out: every current member, ids included."""
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="3F3F76")
    for column, (heading, _) in enumerate(COLUMNS, 1):
        cell = sheet.cell(row=1, column=column, value=heading)
        cell.font, cell.fill = head, fill
        cell.alignment = Alignment(horizontal="center")
        sheet.column_dimensions[get_column_letter(column)].width = WIDTHS[heading]
    for line, player in enumerate(sorted(players, key=lambda p: (-(p.bgb_cp or 0), p.name)), 2):
        for column, (_, attribute) in enumerate(COLUMNS, 1):
            value = str(player.id) if attribute == "id" else getattr(player, attribute)
            cell = sheet.cell(row=line, column=column, value=value)
            if attribute == "id":
                # Not protection, just a hint: this column is the site's, not yours.
                cell.font = Font(color="808080")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(1, len(players) + 1)}"

    help_sheet = book.create_sheet(HELP_SHEET)
    help_sheet.column_dimensions["A"].width = 84
    lines = list(HELP)
    if as_of:
        lines += ["", f"The numbers in this file are what the site held on {as_of}."]
    for line, text in enumerate(lines, 1):
        cell = help_sheet.cell(row=line, column=1, value=text)
        if line == 1:
            cell.font = Font(bold=True, size=13)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _number(value, field_name: str, line: int, problems: list[str]) -> int | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    text = str(value).replace(",", "").replace(" ", "").strip()
    try:
        number = int(float(text))
    except ValueError:
        problems.append(f"Row {line}: “{value}” is not a number, in {field_name}.")
        return None
    low, high = NUMERIC[field_name]
    if not low <= number <= high:
        problems.append(f"Row {line}: {field_name} of {number} is out of range ({low}-{high}).")
        return None
    return number


def read_workbook(data: bytes) -> tuple[list[SheetRow], list[str]]:
    """The rows of an uploaded file, and everything wrong with it."""
    problems: list[str] = []
    try:
        book = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:  # noqa: BLE001 - openpyxl raises several unrelated types
        return [], ["That file could not be opened as a spreadsheet. Save it as .xlsx."]
    sheet = book[SHEET] if SHEET in book.sheetnames else book.worksheets[0]

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [], ["The sheet is empty."]
    headings = {str(v).strip().casefold(): i for i, v in enumerate(rows[0]) if v is not None}
    where = {}
    for heading, attribute in COLUMNS:
        if heading.casefold() not in headings:
            problems.append(f"The column “{heading}” is missing. Download the template again.")
        else:
            where[attribute] = headings[heading.casefold()]
    if problems:
        return [], problems

    out: list[SheetRow] = []
    for line, values in enumerate(rows[1:], 2):
        def cell(attribute: str):
            index = where[attribute]
            return values[index] if index < len(values) else None

        if not any(v not in (None, "") for v in values):
            continue
        name = str(cell("name") or "").strip()
        if not name:
            problems.append(f"Row {line}: no name.")
            continue
        if len(name) > 100:
            problems.append(f"Row {line}: the name is longer than 100 characters.")
            continue
        raw_id = str(cell("id") or "").strip()
        player_id = None
        if raw_id:
            try:
                player_id = uuid.UUID(raw_id)
            except ValueError:
                problems.append(f"Row {line}: “{raw_id}” is not a player id. Leave it empty "
                                "for someone new, or paste the id back from the template.")
                continue
        out.append(SheetRow(
            line=line, id=player_id, name=name,
            rank=_number(cell("rank"), "rank", line, problems),
            industry_level=_number(cell("industry_level"), "industry_level", line, problems),
            bgb_cp=_number(cell("bgb_cp"), "bgb_cp", line, problems),
            total_cp=_number(cell("total_cp"), "total_cp", line, problems),
        ))

    seen_ids: dict[uuid.UUID, int] = {}
    for row in out:
        if row.id and row.id in seen_ids:
            problems.append(f"Row {row.line}: the same player id is also on "
                            f"row {seen_ids[row.id]}.")
        elif row.id:
            seen_ids[row.id] = row.line
    seen_names: dict[str, int] = {}
    for row in out:
        if row.name in seen_names:
            problems.append(f"Row {row.line}: “{row.name}” is also the name on row "
                            f"{seen_names[row.name]}. Two members cannot share a name.")
        else:
            seen_names[row.name] = row.line
    return out, problems


def plan_changes(rows: list[SheetRow], players: list, as_of: date,
                 keep: set[str] | None = None) -> Plan:
    """What the upload would do. `keep` holds ids whose departure you unticked."""
    plan = Plan(as_of=as_of)
    if not rows:
        # Without rows there is nothing to compare against, and every member
        # would look as though they had left. An unreadable file must not be
        # able to empty the roster.
        plan.problems.append("No rows could be read from that sheet, so nothing was worked out.")
        return plan
    by_id = {str(p.id): p for p in players}
    for row in rows:
        if row.id is None:
            plan.added.append(Change(player_id=None, name=row.name, line=row.line, fields={
                f: (None, getattr(row, f)) for f in ("rank", "industry_level", "bgb_cp", "total_cp")
                if getattr(row, f) is not None}))
            continue
        player = by_id.get(str(row.id))
        if player is None:
            plan.problems.append(f"Row {row.line}: no player has the id {row.id}. Leave the id "
                                 "empty if they are new.")
            continue
        fields = {}
        for f in ("rank", "industry_level", "bgb_cp", "total_cp"):
            value = getattr(row, f)
            if value is not None and value != getattr(player, f):
                fields[f] = (getattr(player, f), value)
        change = Change(player_id=str(player.id), name=row.name, fields=fields, line=row.line)
        if not player.active:
            change.was = player.name if player.name != row.name else None
            plan.returning.append(change)
        elif player.name != row.name:
            change.was = player.name
            plan.renamed.append(change)
        elif fields:
            plan.updated.append(change)
        else:
            plan.unchanged += 1

    in_file = {str(row.id) for row in rows if row.id}
    for player in players:
        kept = keep or set()
        if player.active and str(player.id) not in in_file and str(player.id) not in kept:
            plan.left.append(Change(player_id=str(player.id), name=player.name))
    return plan
