"""The two BGB spreadsheets: the roster before a battle, the result after it.

Neither is ours to decide. An admin reads the roster off the "Participants"
popup once registration locks on the Thursday, and the result off the mail that
arrives when the battle ends. All this file does is let them record both
against seats the site already knows, by handing out a sheet per team with the
ids filled in and the columns to fill blank.

Ids are what make it reliable, exactly as in `roster_sheet`: a mark belongs to a
row, and the row names an id, so a member who renamed between the template and
the upload still lands in the right seat. The roster sheet carries player ids;
the result sheet carries registration ids, because a mercenary is nobody's
player and still has to be scored.

Nothing here touches the database. It produces a plan; the endpoint applies it
once the person uploading has seen it.
"""
from __future__ import annotations

import hashlib
import io
import uuid
from dataclasses import dataclass, field
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TEAMS = ("A", "B")
SHEETS = {"A": "Team A", "B": "Team B"}
HELP_SHEET = "How to use"

# What the game allows per team. Registration cannot exceed these, so a sheet
# that does was misread rather than unusual.
MAX_STARTERS = 20
MAX_SUBSTITUTES = 10

STARTER, SUBSTITUTE = "starter", "substitute"
MERCENARY = "mercenary"

COLUMNS: list[tuple[str, str]] = [
    ("Player id", "id"),
    ("Name", "name"),
    ("BGB CP", "bgb_cp"),
    ("Starter (O/X)", STARTER),
    ("Substitute (O/X)", SUBSTITUTE),
    ("Mercenary (O/X)", MERCENARY),
]
MARKS = (STARTER, SUBSTITUTE, MERCENARY)
WIDTHS = {"Player id": 38, "Name": 24, "BGB CP": 16,
          "Starter (O/X)": 14, "Substitute (O/X)": 17, "Mercenary (O/X)": 16}

# Typed by hand in a hurry, so read generously — but only where the meaning is
# beyond doubt. Anything else is reported rather than guessed at.
YES = {"O", "0", "○", "◯", "●", "V", "Y", "YES", "✓", "✔", "TRUE"}
NO = {"", "X", "×", "✗", "-", "–", "N", "NO", "FALSE"}

HELP = [
    "How to record a BGB registration",
    "",
    "1. Each sheet is one team. Every current member is listed on both, because",
    "   the game lets you put anyone on either side.",
    "2. Once registration locks, read the in-game Participants popup and mark the",
    "   Starter or Substitute column with O. Leave everyone else blank, or X.",
    "3. Upload the file on the BGB tab. The site shows the roster it read and",
    "   asks you to confirm before recording anything.",
    "",
    f"A team takes at most {MAX_STARTERS} starters and {MAX_SUBSTITUTES} substitutes.",
    "Nobody can be on both teams, and nobody can be both starter and substitute.",
    "",
    "The Player id column is the site's, not yours — leave it alone. A member who",
    "is missing from these sheets has to be added on the Members tab first, since",
    "a registration is recorded against a player the site already knows.",
    "",
    "Mercenaries from another alliance are the exception. Add a row at the bottom,",
    "leave Player id empty, type their name and BGB CP, mark O under Mercenary, and",
    "mark Starter or Substitute as usual. They count towards the team's limits and",
    "take a seat on the card like anyone else. They are recorded against this",
    "battle only — they never join the Members tab, because the member import reads",
    "an absent row as somebody who left, and they were never here to leave.",
    "",
    "BGB CP is what the site holds today. Correcting a number here updates the",
    "member's BGB CP as well, so the lineup cards are drafted on the real figures.",
    "A mercenary's CP is only ever used for this battle.",
]


@dataclass
class SheetRow:
    """One row as it came back, already checked."""

    team: str
    line: int
    id: uuid.UUID | None               # None only for a mercenary, who is nobody's
    name: str
    bgb_cp: int | None = None
    role: str | None = None            # starter | substitute | None (not registered)
    mercenary: bool = False


@dataclass
class Seat:
    """Somebody registered for a battle: a member, or a mercenary hired for it."""

    player_id: str | None              # None for a mercenary
    name: str
    team: str
    role: str
    bgb_cp: int | None = None
    mercenary: bool = False


@dataclass
class CpChange:
    player_id: str
    name: str
    before: int | None
    after: int


@dataclass
class Plan:
    seats: list[Seat] = field(default_factory=list)
    cp_changes: list[CpChange] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)

    def of(self, team: str, role: str) -> list[Seat]:
        return [s for s in self.seats if s.team == team and s.role == role]


def fingerprint(players: list) -> str:
    """A short digest of what the template was built from, to notice it moving."""
    parts = sorted(f"{p.id}|{p.name}|{p.bgb_cp}|{p.active}" for p in players)
    return hashlib.sha256("\n".join(parts).encode()).hexdigest()[:32]


def build_template(players: list, as_of: date | None = None) -> bytes:
    """The workbook to hand out: one sheet per team, every member on both."""
    book = Workbook()
    book.remove(book.active)
    ordered = sorted(players, key=lambda p: (-(p.bgb_cp or 0), p.name))
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="3F3F76")
    mark = PatternFill("solid", fgColor="FFF2CC")

    for team in TEAMS:
        sheet = book.create_sheet(SHEETS[team])
        for column, (heading, _) in enumerate(COLUMNS, 1):
            cell = sheet.cell(row=1, column=column, value=heading)
            cell.font, cell.fill = head, fill
            cell.alignment = Alignment(horizontal="center")
            sheet.column_dimensions[get_column_letter(column)].width = WIDTHS[heading]
        for line, player in enumerate(ordered, 2):
            for column, (_, attribute) in enumerate(COLUMNS, 1):
                if attribute in MARKS:
                    # The columns to fill in, tinted so they are obvious.
                    cell = sheet.cell(row=line, column=column)
                    cell.fill = mark
                    cell.alignment = Alignment(horizontal="center")
                    continue
                value = str(player.id) if attribute == "id" else getattr(player, attribute)
                cell = sheet.cell(row=line, column=column, value=value)
                if attribute == "id":
                    # Not protection, just a hint: this column is the site's, not yours.
                    cell.font = Font(color="808080")
        # A mercenary has no row of their own until one is typed, so the sheet
        # says where to type it rather than leaving the space unexplained.
        note = sheet.cell(row=len(ordered) + 3, column=1,
                          value="Mercenaries from another alliance: add rows here. "
                                "Leave Player id empty, fill in Name and BGB CP, and mark "
                                "O under Mercenary as well as Starter or Substitute.")
        note.font = Font(color="808080", italic=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(COLUMNS))}{max(1, len(ordered) + 1)}"
        )

    help_sheet = book.create_sheet(HELP_SHEET)
    help_sheet.column_dimensions["A"].width = 84
    lines = list(HELP)
    if as_of:
        lines += ["", f"The CP in this file is what the site held on {as_of}."]
    for line, text in enumerate(lines, 1):
        cell = help_sheet.cell(row=line, column=1, value=text)
        if line == 1:
            cell.font = Font(bold=True, size=13)

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _mark(value, heading: str, team: str, line: int, problems: list[str]) -> bool | None:
    text = str(value if value is not None else "").strip().upper()
    if text in YES:
        return True
    if text in NO:
        return False
    problems.append(f"{SHEETS[team]} row {line}: “{value}” is not an O or an X, "
                    f"in {heading}.")
    return None


def _number(value, team: str, line: int, problems: list[str]) -> int | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    text = str(value).replace(",", "").replace(" ", "").strip()
    try:
        number = int(float(text))
    except ValueError:
        problems.append(f"{SHEETS[team]} row {line}: “{value}” is not a number, in BGB CP.")
        return None
    if not 0 <= number <= 10**12:
        problems.append(f"{SHEETS[team]} row {line}: a BGB CP of {number} is out of range.")
        return None
    return number


def read_workbook(data: bytes) -> tuple[list[SheetRow], list[str]]:
    """The marked rows of an uploaded file, and everything wrong with it."""
    problems: list[str] = []
    try:
        book = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:  # noqa: BLE001 - openpyxl raises several unrelated types
        return [], ["That file could not be opened as a spreadsheet. Save it as .xlsx."]

    missing = [SHEETS[t] for t in TEAMS if SHEETS[t] not in book.sheetnames]
    if missing:
        return [], [f"The sheet “{missing[0]}” is missing. Download the template again."]

    out: list[SheetRow] = []
    for team in TEAMS:
        sheet = book[SHEETS[team]]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            problems.append(f"{SHEETS[team]} is empty.")
            continue
        headings = {str(v).strip().casefold(): i for i, v in enumerate(rows[0]) if v is not None}
        where = {}
        for heading, attribute in COLUMNS:
            if heading.casefold() not in headings:
                problems.append(f"{SHEETS[team]}: the column “{heading}” is missing. "
                                "Download the template again.")
            else:
                where[attribute] = headings[heading.casefold()]
        if len(where) < len(COLUMNS):
            continue

        seen: dict[uuid.UUID, int] = {}
        hired: dict[str, int] = {}
        for line, values in enumerate(rows[1:], 2):
            def cell(attribute: str, values=values, where=where):
                index = where[attribute]
                return values[index] if index < len(values) else None

            if not any(v not in (None, "") for v in values):
                continue
            raw_id = str(cell("id") or "").strip()
            name = str(cell("name") or "").strip()
            starter = _mark(cell(STARTER), "Starter", team, line, problems)
            substitute = _mark(cell(SUBSTITUTE), "Substitute", team, line, problems)
            mercenary = _mark(cell(MERCENARY), "Mercenary", team, line, problems)
            try:
                player_id = uuid.UUID(raw_id) if raw_id else None
            except ValueError:
                player_id = None
            # Nothing marked and no id to match anyone by: not a registration at
            # all. The note at the foot of the sheet reads like this, and so does
            # anything jotted in the margin.
            if player_id is None and not (starter or substitute or mercenary):
                continue
            if starter and substitute:
                problems.append(f"{SHEETS[team]} row {line}: “{name}” is marked as both a "
                                "starter and a substitute.")
                continue
            role = STARTER if starter else SUBSTITUTE if substitute else None
            bgb_cp = _number(cell("bgb_cp"), team, line, problems)

            if mercenary:
                # Hired for this battle from another alliance. They have no member
                # record, so everything the card needs has to be on the row.
                if raw_id:
                    problems.append(f"{SHEETS[team]} row {line}: “{name}” has a player id, so "
                                    "they are a member. Clear the Mercenary column, or clear "
                                    "the id if this is somebody else.")
                    continue
                if not name:
                    problems.append(f"{SHEETS[team]} row {line}: a mercenary needs a name.")
                    continue
                if bgb_cp is None:
                    problems.append(f"{SHEETS[team]} row {line}: “{name}” needs a BGB CP. "
                                    "It is what decides which seat they take.")
                    continue
                if role is None:
                    problems.append(f"{SHEETS[team]} row {line}: “{name}” is marked as a "
                                    "mercenary but as neither a starter nor a substitute.")
                    continue
                if name in hired:
                    problems.append(f"{SHEETS[team]} row {line}: “{name}” is also hired on "
                                    f"row {hired[name]}.")
                    continue
                hired[name] = line
                out.append(SheetRow(team=team, line=line, id=None, name=name, role=role,
                                    bgb_cp=bgb_cp, mercenary=True))
                continue

            if not raw_id:
                problems.append(
                    f"{SHEETS[team]} row {line}: no player id"
                    + (f" for “{name}”" if name else "")
                    + ". Add them on the Members tab first and download the template again, "
                      "or mark O under Mercenary if they are from another alliance.")
                continue
            if player_id is None:
                problems.append(f"{SHEETS[team]} row {line}: “{raw_id}” is not a player id.")
                continue
            if player_id in seen:
                problems.append(f"{SHEETS[team]} row {line}: the same player is also on "
                                f"row {seen[player_id]}.")
                continue
            seen[player_id] = line
            out.append(SheetRow(
                team=team, line=line, id=player_id, name=name, role=role, bgb_cp=bgb_cp,
            ))
    return out, problems


def plan_roster(rows: list[SheetRow], players: list) -> Plan:
    """What the upload would record. Reads nothing, writes nothing."""
    plan = Plan()
    if not rows:
        # Same guard as the member import: a file that failed to parse must not
        # be able to record an empty roster over a real one.
        plan.problems.append("No rows could be read from that sheet, so nothing was worked out.")
        return plan

    by_id = {str(p.id): p for p in players}
    marked = [r for r in rows if r.role]
    if not marked:
        plan.problems.append("Nobody is marked as a starter or a substitute, so there is "
                             "no roster to record.")
        return plan

    current = {p.name for p in players if p.active}
    cp_seen: dict[str, CpChange] = {}
    for row in rows:
        if row.mercenary:
            # Nobody's member, so there is no record to check them against and
            # nothing of theirs to update -- except that they must not be one of
            # ours under a hand-typed name, which would put them on twice.
            if row.name in current:
                plan.problems.append(f"{SHEETS[row.team]} row {row.line}: “{row.name}” is "
                                     "already a member. Mark their own row instead.")
                continue
            plan.seats.append(Seat(player_id=None, name=row.name, team=row.team,
                                   role=row.role, bgb_cp=row.bgb_cp, mercenary=True))
            continue
        player = by_id.get(str(row.id))
        if player is None:
            if row.role:
                plan.problems.append(f"{SHEETS[row.team]} row {row.line}: no member has the "
                                     f"id {row.id}.")
            continue
        # A corrected CP updates the member as well, so the cards are drafted on
        # the real figure -- but only ever once, however many sheets carry it.
        if row.bgb_cp is not None and row.bgb_cp != player.bgb_cp:
            cp_seen.setdefault(str(player.id), CpChange(
                player_id=str(player.id), name=player.name,
                before=player.bgb_cp, after=row.bgb_cp))
        if not row.role:
            continue
        if not player.active:
            plan.problems.append(f"{SHEETS[row.team]} row {row.line}: “{player.name}” has left "
                                 "the alliance. Bring them back on the Members tab first.")
            continue
        plan.seats.append(Seat(
            player_id=str(player.id), name=player.name, team=row.team, role=row.role,
            bgb_cp=row.bgb_cp if row.bgb_cp is not None else player.bgb_cp))
    plan.cp_changes = list(cp_seen.values())

    both: dict[str, str] = {}
    for seat in plan.seats:
        # A mercenary has no id, so they are told apart by name -- which is all
        # we know them by, and why two of them may not share one.
        who = seat.player_id or f"hired:{seat.name}"
        if who in both and both[who] != seat.team:
            plan.problems.append(f"“{seat.name}” is registered on both teams. Nobody plays "
                                 "for both sides.")
        both[who] = seat.team

    for team in TEAMS:
        for role, cap, word in ((STARTER, MAX_STARTERS, "starters"),
                                (SUBSTITUTE, MAX_SUBSTITUTES, "substitutes")):
            count = len(plan.of(team, role))
            if count > cap:
                plan.problems.append(f"{SHEETS[team]} has {count} {word} marked, and the game "
                                     f"allows {cap}.")
    return plan


# ---------------------------------------------------------------------------
# The result: the same roster, read back with what everyone scored.
# ---------------------------------------------------------------------------
#
# The game's result mail is one ranking of every participant from both
# alliances, so the only rows that matter are ours -- which the roster already
# names. The sheet handed out is therefore the battle's own seats, with a
# column to fill in per seat, and each row carries its registration id rather
# than a player id, because a mercenary has no player to point at.

PARTICIPATED = "participated"

RESULT_COLUMNS: list[tuple[str, str]] = [
    ("Registration id", "id"),
    ("Name", "name"),
    ("Team", "team"),
    ("Role", "role"),
    ("BGB CP", "bgb_cp"),
    ("Score", "score"),
    ("Participated (O/X)", PARTICIPATED),
]
RESULT_WIDTHS = {"Registration id": 15, "Name": 24, "Team": 8, "Role": 12,
                 "BGB CP": 16, "Score": 14, "Participated (O/X)": 18}
# Filled in from the screenshots; everything else is there to read against them.
RESULT_FILLED = {"score", PARTICIPATED}

RESULT_HELP = [
    "How to record a BGB result",
    "",
    "1. Each sheet is one team, and lists everyone who was registered for that",
    "   battle -- starters, substitutes and mercenaries alike. Nobody else can",
    "   appear in the result, so there are no rows to add.",
    "2. Open the battle-result mail in the game. It ranks every participant from",
    "   both alliances together, so read only the rows tagged with our own",
    "   alliance and ignore the opponent's.",
    "3. Type each player's Score exactly as the game prints it: 2M, 1.6M, 993.7K,",
    "   651.6K, or a plain number. Then upload the file on the BGB tab.",
    "",
    "The two ways of not fighting are told apart by what you type:",
    "",
    "  * Score 0 -- the game listed them with a zero. They were registered and",
    "    never fought.",
    "  * Score left empty -- they are not in the ranking at all, which means they",
    "    were dropped from the line-up before the battle began.",
    "",
    "Participated is worked out from the score, so leave it empty unless you",
    "disagree with it. Mark O or X only where you know better than the number,",
    "and the site will say so rather than quietly picking one.",
    "",
    "A starter who did not fight is the one thing worth chasing, so the site",
    "lists them on their own once the file is read.",
]


@dataclass
class ResultRow:
    """One seat as its result came back."""

    line: int
    id: int
    score: int | None = None
    participated: bool | None = None       # None = say nothing, take the score's word


@dataclass
class Outcome:
    """What one seat did, once the score and the mark have been read together."""

    registration_id: int
    name: str
    team: str
    role: str
    score: int | None
    participated: bool
    listed: bool                           # in the game's ranking at all
    bgb_cp: int | None = None

    @property
    def no_show(self) -> bool:
        """A starter who did not fight: the one thing worth following up."""
        return self.role == STARTER and not self.participated


@dataclass
class ResultPlan:
    outcomes: list[Outcome] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)

    def of(self, team: str) -> list[Outcome]:
        return [o for o in self.outcomes if o.team == team]

    @property
    def no_shows(self) -> list[Outcome]:
        return [o for o in self.outcomes if o.no_show]


def result_fingerprint(registrations: list) -> str:
    """A digest of the seats, to notice the roster moving under an upload."""
    parts = sorted(f"{r.id}|{r.name}|{r.team}|{r.role}" for r in registrations)
    return hashlib.sha256("\n".join(parts).encode()).hexdigest()[:32]


def build_result_template(registrations: list, battle_date: date) -> bytes:
    """The workbook to hand out: this battle's seats, with a column to fill in."""
    book = Workbook()
    book.remove(book.active)
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="3F3F76")
    mark = PatternFill("solid", fgColor="FFF2CC")
    order = {STARTER: 0, SUBSTITUTE: 1}

    for team in TEAMS:
        seats = sorted((r for r in registrations if r.team == team),
                       key=lambda r: (order.get(r.role, 9), -(r.bgb_cp or 0), r.name))
        sheet = book.create_sheet(SHEETS[team])
        for column, (heading, _) in enumerate(RESULT_COLUMNS, 1):
            cell = sheet.cell(row=1, column=column, value=heading)
            cell.font, cell.fill = head, fill
            cell.alignment = Alignment(horizontal="center")
            sheet.column_dimensions[get_column_letter(column)].width = RESULT_WIDTHS[heading]
        for line, seat in enumerate(seats, 2):
            for column, (_, attribute) in enumerate(RESULT_COLUMNS, 1):
                if attribute in RESULT_FILLED:
                    cell = sheet.cell(row=line, column=column)
                    cell.fill = mark
                    cell.alignment = Alignment(horizontal="center")
                    continue
                value = getattr(seat, attribute)
                cell = sheet.cell(row=line, column=column, value=value)
                if attribute == "id":
                    # Not protection, just a hint: this column is the site's.
                    cell.font = Font(color="808080")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(RESULT_COLUMNS))}{max(1, len(seats) + 1)}"
        )

    help_sheet = book.create_sheet(HELP_SHEET)
    help_sheet.column_dimensions["A"].width = 84
    for line, text in enumerate([*RESULT_HELP, "", f"This is the battle of {battle_date}."], 1):
        cell = help_sheet.cell(row=line, column=1, value=text)
        if line == 1:
            cell.font = Font(bold=True, size=13)

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


# The game prints a score the way it prints CP: 2M, 1.6M, 993.7K, 651.6K. Typed
# straight off the screen, so read straight off the screen.
SUFFIXES = {"K": 1_000, "M": 1_000_000, "B": 1_000_000_000}


def _score(value, team: str, line: int, problems: list[str]) -> int | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    text = str(value).replace(",", "").replace(" ", "").strip().upper()
    multiplier = 1
    if text and text[-1] in SUFFIXES:
        multiplier = SUFFIXES[text[-1]]
        text = text[:-1]
    try:
        number = int(round(float(text) * multiplier))
    except ValueError:
        problems.append(f"{SHEETS[team]} row {line}: “{value}” is not a score. Type it the "
                        "way the game prints it — 2M, 993.7K — or as a plain number.")
        return None
    if not 0 <= number <= 10**12:
        problems.append(f"{SHEETS[team]} row {line}: a score of {number} is out of range.")
        return None
    return number


def _stated(value, team: str, line: int, problems: list[str]) -> bool | None:
    """The Participated mark, where an empty cell says nothing rather than no.

    The roster sheet reads a blank as an X, because most of its rows are blank.
    Here the column is the exception, only filled where the admin disagrees with
    the score — so a blank has to mean "take the score's word", or every row
    they did not tick would contradict the number beside it.
    """
    text = str(value if value is not None else "").strip().upper()
    if not text:
        return None
    if text in YES:
        return True
    if text in NO:
        return False
    problems.append(f"{SHEETS[team]} row {line}: “{value}” is not an O or an X, "
                    "in Participated.")
    return None


def read_result_workbook(data: bytes) -> tuple[list[ResultRow], list[str]]:
    """The filled-in rows of an uploaded result sheet, and what is wrong with it."""
    problems: list[str] = []
    try:
        book = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:  # noqa: BLE001 - openpyxl raises several unrelated types
        return [], ["That file could not be opened as a spreadsheet. Save it as .xlsx."]

    missing = [SHEETS[t] for t in TEAMS if SHEETS[t] not in book.sheetnames]
    if missing:
        return [], [f"The sheet “{missing[0]}” is missing. Download the template again."]

    out: list[ResultRow] = []
    for team in TEAMS:
        sheet = book[SHEETS[team]]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        headings = {str(v).strip().casefold(): i for i, v in enumerate(rows[0]) if v is not None}
        where = {}
        for heading, attribute in RESULT_COLUMNS:
            if heading.casefold() not in headings:
                problems.append(f"{SHEETS[team]}: the column “{heading}” is missing. "
                                "Download the template again.")
            else:
                where[attribute] = headings[heading.casefold()]
        if len(where) < len(RESULT_COLUMNS):
            continue

        seen: dict[int, int] = {}
        for line, values in enumerate(rows[1:], 2):
            def cell(attribute: str, values=values, where=where):
                index = where[attribute]
                return values[index] if index < len(values) else None

            if not any(v not in (None, "") for v in values):
                continue
            raw_id = str(cell("id") or "").strip()
            if not raw_id:
                problems.append(f"{SHEETS[team]} row {line}: no registration id. The result is "
                                "recorded against the roster, so rows cannot be added here — "
                                "fix the roster first if somebody is missing.")
                continue
            try:
                seat_id = int(float(raw_id))
            except ValueError:
                problems.append(f"{SHEETS[team]} row {line}: “{raw_id}” is not a registration id.")
                continue
            if seat_id in seen:
                problems.append(f"{SHEETS[team]} row {line}: the same seat is also on "
                                f"row {seen[seat_id]}.")
                continue
            seen[seat_id] = line
            out.append(ResultRow(
                line=line, id=seat_id,
                score=_score(cell("score"), team, line, problems),
                participated=_stated(cell(PARTICIPATED), team, line, problems),
            ))
    return out, problems


def plan_result(rows: list[ResultRow], registrations: list) -> ResultPlan:
    """What the upload would record against a battle's seats."""
    plan = ResultPlan()
    if not rows:
        # The roster's own guard, for the same reason: a file that failed to
        # parse must not be able to mark a whole alliance as having not fought.
        plan.problems.append("No rows could be read from that sheet, so nothing was worked out.")
        return plan

    by_id = {r.id: r for r in registrations}
    covered = set()
    for row in rows:
        seat = by_id.get(row.id)
        if seat is None:
            plan.problems.append(f"Row {row.line}: no seat {row.id} in this battle's roster. "
                                 "The sheet may be for a different date.")
            continue
        covered.add(seat.id)
        listed = row.score is not None
        participated = row.participated if row.participated is not None else bool(row.score)
        # A mark that disagrees with the number is reported, never resolved
        # quietly: one of the two was read wrong, and only the admin knows which.
        if row.participated is True and not row.score:
            plan.problems.append(
                f"{SHEETS[seat.team]} row {row.line}: “{seat.name}” is marked as having "
                + ("fought, but scored zero." if listed else
                   "fought, but has no score. Leave Participated empty, or give them a score."))
            continue
        if row.participated is False and row.score:
            plan.problems.append(f"{SHEETS[seat.team]} row {row.line}: “{seat.name}” scored "
                                 f"{row.score:,} but is marked as not having fought.")
            continue
        plan.outcomes.append(Outcome(
            registration_id=seat.id, name=seat.name, team=seat.team, role=seat.role,
            score=row.score, participated=participated, listed=listed, bgb_cp=seat.bgb_cp))

    absent = [r for r in registrations if r.id not in covered]
    if absent:
        plan.problems.append(
            f"{len(absent)} of the battle's {len(registrations)} seats are not in the file "
            f"({', '.join(r.name for r in absent[:3])}"
            f"{'...' if len(absent) > 3 else ''}). Download the template again — a result is "
            "recorded for the whole roster at once.")
    return plan
